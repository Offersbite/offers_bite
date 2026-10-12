import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

from app.services import offer_import


def build_workbook() -> bytes:
    workbook = Workbook()
    cards = workbook.active
    cards.title = "Card offers"
    cards.append([
        "Bank Name",
        "Card Name",
        "Card Type",
        "Merchant",
        "Category",
        "Offer Headline",
        "Max Discount (INR)",
        "Min purchase",
        "Valid Days",
        "Expiry Date",
        "Usage Limit",
    ])
    cards.append([
        "Example Bank",
        "Rewards Visa Credit Card",
        "Credit Card",
        "Example Shop",
        "Shopping",
        "Save INR 100",
        100,
        500,
        "Friday to Sunday",
        "2027-12-31",
        "Once per customer",
    ])
    cards.append([
        "Example Bank",
        "Rewards Debit Card",
        "Debit",
        "Example Shop",
        "Shopping",
        "No expiry offer",
        None,
        None,
        "All days",
        "N/A",
        "No limit",
    ])
    payments = workbook.create_sheet("Payment offers")
    payments.append(["Bank / Provider", "Card Name", "Card Type", "Merchant", "Offer Headline", "Valid Days", "Expiry Date"])
    payments.append(["Example Wallet", "Example UPI", "UPI", "Example Shop", "UPI cashback", "Any", "N/A"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class OfferWorkbookParserTests(unittest.TestCase):
    def test_parses_all_sheets_and_retains_original_columns(self) -> None:
        rows = offer_import.parse_offer_workbook(build_workbook())

        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0].source_sheet, "Card offers")
        self.assertEqual(rows[0].source_row, 2)
        self.assertEqual(rows[0].valid_days, [0, 5, 6])
        self.assertEqual(rows[0].source_data["Usage Limit"], "Once per customer")
        self.assertTrue(rows[0].is_matchable)
        self.assertIsNone(rows[1].expires_on)
        self.assertEqual(rows[1].max_discount, None)
        self.assertFalse(rows[2].is_matchable)
        self.assertEqual(rows[2].unmatched_reason, "Payment type is not a supported bank card.")

    def test_rejects_bad_headers_before_import(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Merchant", "Wrong column"])
        sheet.append(["Shop", "Offer"])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()

        with self.assertRaises(offer_import.OfferWorkbookError) as result:
            offer_import.parse_offer_workbook(stream.getvalue())

        self.assertIn("Merchant and Offer Headline", result.exception.errors[0]["message"])

    def test_rejects_corrupt_workbook(self) -> None:
        with self.assertRaises(offer_import.OfferWorkbookError):
            offer_import.parse_offer_workbook(b"not an xlsx")


class OfferWorkbookImportTests(unittest.TestCase):
    def test_imports_card_rows_and_preserves_non_card_rows_without_matching(self) -> None:
        rows = offer_import.parse_offer_workbook(build_workbook())
        master_cards = MagicMock()
        master_cards.find.return_value = []
        merchant_offers = MagicMock()
        merchant_offers.find.side_effect = [
            [],
            [
                {"source_sheet": row.source_sheet, "source_row": row.source_row}
                for row in rows
            ],
        ]
        merchant_offers.insert_many.return_value.inserted_ids = ["one", "two", "three"]

        with (
            patch.object(offer_import, "master_cards", master_cards),
            patch.object(offer_import, "merchant_offers", merchant_offers),
            patch.object(offer_import, "new_catalogue_id", side_effect=[101, 102, 201, 202, 203]),
        ):
            result = offer_import.import_offer_rows(
                rows,
                source_file="test.xlsx",
                source_hash="f" * 64,
            )

        self.assertEqual(result["rows_imported"], 3)
        self.assertEqual(result["cards_created"], 2)
        self.assertEqual(result["unmatched_rows"], 1)
        inserted = merchant_offers.insert_many.call_args.args[0]
        self.assertEqual(inserted[0]["source_data"]["Usage Limit"], "Once per customer")
        self.assertEqual(inserted[0]["valid_days"], [0, 5, 6])
        self.assertEqual(inserted[2]["applicable_master_card_ids"], [])
        self.assertFalse(inserted[2]["is_matchable"])
        self.assertIsNone(inserted[1]["expires_on"])


if __name__ == "__main__":
    unittest.main()
