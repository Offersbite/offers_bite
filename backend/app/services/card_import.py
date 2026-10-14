import csv
from io import BytesIO, StringIO
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.schemas.admin import MasterCardData

MAX_CARD_ROWS = 10_000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_WORKBOOK_ENTRIES = 5_000
CARD_TYPES = {"CREDIT", "DEBIT"}
CARD_NETWORKS = {"VISA", "MASTERCARD", "RUPAY", "AMEX", "DINERS", "OTHER"}
REQUIRED_FIELDS = {"bank_name", "card_name", "card_type", "network"}

HEADER_ALIASES = {
    "bank": "bank_name",
    "bank_name": "bank_name",
    "bankname": "bank_name",
    "card": "card_name",
    "card_name": "card_name",
    "cardname": "card_name",
    "card_typ": "card_type",
    "card_type": "card_type",
    "cardtype": "card_type",
    "type": "card_type",
    "network": "network",
    "card_network": "network",
    "cardnetwork": "network",
    "tier": "tier_category",
    "tier_category": "tier_category",
    "tiercategory": "tier_category",
    "tier_name": "tier_category",
    "category_tier": "tier_category",
}


class CardImportError(ValueError):
    def __init__(self, message: str, errors: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.errors = errors or []


def _normalise_header(value: Any) -> str:
    header = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return HEADER_ALIASES.get(header, header)


def _normalise_card_type(value: Any) -> str:
    raw = str(value or "").strip().upper()
    if "DEBIT" in raw:
        return "DEBIT"
    if "CREDIT" in raw:
        return "CREDIT"
    return raw


def _normalise_network(value: Any) -> str:
    raw = str(value or "").strip().upper()
    if "VISA" in raw:
        return "VISA"
    if "MASTER" in raw:
        return "MASTERCARD"
    if "RUPAY" in raw:
        return "RUPAY"
    if "AMEX" in raw or "AMERICAN EXPRESS" in raw:
        return "AMEX"
    if "DINER" in raw:
        return "DINERS"
    return raw if raw in CARD_NETWORKS else "OTHER"


def _validate_row_dicts(
    rows: list[tuple[str, int, dict[str, Any]]],
) -> list[dict[str, str]]:
    imported: list[dict[str, str]] = []
    errors: list[dict[str, Any]] = []
    processed = 0

    for sheet_name, source_row, raw in rows:
        if not any(value is not None and str(value).strip() for value in raw.values()):
            continue
        processed += 1
        if processed > MAX_CARD_ROWS:
            errors.append({
                "sheet": sheet_name,
                "row": source_row,
                "message": f"The file exceeds the {MAX_CARD_ROWS:,}-row limit.",
            })
            break

        bank_name = str(raw.get("bank_name") or "").strip()
        card_name = str(raw.get("card_name") or "").strip()
        card_type = _normalise_card_type(raw.get("card_type"))
        network = _normalise_network(raw.get("network"))
        tier_category = str(raw.get("tier_category") or "").strip().upper() or "STANDARD"

        try:
            card = MasterCardData.model_validate({
                "bank_name": bank_name,
                "card_name": card_name,
                "card_type": card_type,
                "network": network,
                "tier_category": tier_category,
            })
        except ValidationError as exc:
            errors.append({
                "sheet": sheet_name,
                "row": source_row,
                "message": "; ".join(
                    f"{'.'.join(map(str, item['loc']))}: {item['msg']}"
                    for item in exc.errors()
                ),
            })
            continue
        imported.append(card.model_dump())

    if errors:
        raise CardImportError(
            f"File validation failed for {len(errors)} row(s); nothing was imported.",
            errors,
        )
    if not imported:
        raise CardImportError("The file contains no card rows.")
    return imported


def parse_card_csv(content: bytes) -> list[dict[str, str]]:
    if not content:
        raise CardImportError("The selected CSV file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise CardImportError("The file exceeds the 10 MB upload limit.")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CardImportError("The CSV file must use UTF-8 encoding.") from exc

    reader = csv.reader(StringIO(text))
    headers = next(reader, [])
    normalised = [_normalise_header(value) for value in headers]
    if len(normalised) != len(set(normalised)):
        raise CardImportError("The CSV file contains duplicate column headers.")
    missing = sorted(REQUIRED_FIELDS - set(normalised))
    if missing:
        raise CardImportError(
            "The CSV file is missing required columns: " + ", ".join(missing) + "."
        )

    row_dicts: list[tuple[str, int, dict[str, Any]]] = []
    for row_number, values in enumerate(reader, start=2):
        row_map = {
            normalised[i]: values[i] if i < len(values) else None
            for i in range(len(normalised))
        }
        row_dicts.append(("CSV", row_number, row_map))

    return _validate_row_dicts(row_dicts)


def parse_card_workbook(content: bytes) -> list[dict[str, str]]:
    if not content:
        raise CardImportError("The selected workbook is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise CardImportError("The file exceeds the 10 MB upload limit.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_WORKBOOK_ENTRIES:
                raise CardImportError("The workbook contains too many internal files.")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise CardImportError("The workbook expands beyond the 50 MB safety limit.")
    except BadZipFile as exc:
        raise CardImportError("The selected file is not a valid .xlsx workbook.") from exc

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise CardImportError("The selected file is not a readable .xlsx workbook.") from exc

    all_rows: list[tuple[str, int, dict[str, Any]]] = []
    try:
        if not workbook.worksheets:
            raise CardImportError("The workbook has no worksheets.")
        for sheet in workbook.worksheets:
            headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            normalised = [_normalise_header(value) for value in headers]
            if len(normalised) != len(set(normalised)):
                raise CardImportError(
                    f"Worksheet '{sheet.title}' contains duplicate column headers."
                )
            missing = sorted(REQUIRED_FIELDS - set(normalised))
            if missing:
                raise CardImportError(
                    f"Worksheet '{sheet.title}' is missing required columns: "
                    + ", ".join(missing) + "."
                )
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                row_map = {
                    normalised[i]: values[i] if i < len(values) else None
                    for i in range(len(normalised))
                }
                all_rows.append((sheet.title, row_number, row_map))
    finally:
        workbook.close()
    return _validate_row_dicts(all_rows)
