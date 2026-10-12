from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from io import BytesIO
import math
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pymongo.errors import BulkWriteError, PyMongoError

from app.db.mongo import master_cards, merchant_offers
from app.services.catalogue import new_catalogue_id

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_WORKBOOK_ENTRIES = 5000
MAX_OFFER_ROWS = 10_000

_DAY_NAMES = {
    "sun": 0,
    "sunday": 0,
    "mon": 1,
    "monday": 1,
    "tue": 2,
    "tues": 2,
    "tuesday": 2,
    "wed": 3,
    "wednesday": 3,
    "thu": 4,
    "thur": 4,
    "thurs": 4,
    "thursday": 4,
    "fri": 5,
    "friday": 5,
    "sat": 6,
    "saturday": 6,
}
_ALL_DAYS = list(range(7))


class OfferWorkbookError(ValueError):
    def __init__(
        self,
        message: str,
        errors: list[dict[str, Any]] | None = None,
        *,
        status_code: int = 422,
    ):
        super().__init__(message)
        self.errors = errors or []
        self.status_code = status_code


@dataclass(frozen=True)
class WorkbookOfferRow:
    source_sheet: str
    source_row: int
    source_data: dict[str, Any]
    merchant_name: str
    category: str
    offer_headline: str
    coupon_code: str | None
    min_spend: float
    max_discount: float | None
    valid_days: list[int]
    expires_on: datetime | None
    bank_name: str | None
    card_name: str | None
    card_types: tuple[str, ...]
    is_matchable: bool
    unmatched_reason: str | None


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _source_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _number(value: Any, *, field: str, sheet: str, row: int) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise OfferWorkbookError(
            f"{sheet}, row {row}: {field} must be a number.",
            [{"sheet": sheet, "row": row, "message": f"{field} must be a number."}],
        )
    try:
        if isinstance(value, (int, float)):
            parsed = float(value)
        else:
            cleaned = re.sub(r"[^0-9.+-]", "", str(value).replace(",", ""))
            parsed = float(cleaned)
    except (TypeError, ValueError) as exc:
        raise OfferWorkbookError(
            f"{sheet}, row {row}: {field} must be a number.",
            [{"sheet": sheet, "row": row, "message": f"{field} must be a number."}],
        ) from exc
    if not math.isfinite(parsed) or parsed < 0:
        raise OfferWorkbookError(
            f"{sheet}, row {row}: {field} cannot be negative.",
            [{"sheet": sheet, "row": row, "message": f"{field} cannot be negative."}],
        )
    return parsed


def _parse_expiry(value: Any, *, sheet: str, row: int) -> datetime | None:
    if value is None or _cell_text(value).lower() in {"", "n/a", "na", "none"}:
        return None
    if re.search(r"\b\d+\s+days?\b.*credited cashback", _cell_text(value), re.IGNORECASE):
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        text = _cell_text(value)
        parsed = None
        for date_format in ("%Y-%m-%d", "%b %d, %Y", "%B %d, %Y", "%d-%b-%Y"):
            try:
                parsed = datetime.strptime(text, date_format)
                break
            except ValueError:
                continue
        if parsed is None:
            raise OfferWorkbookError(
                f"{sheet}, row {row}: expiry date {text!r} is not recognised.",
                [{"sheet": sheet, "row": row, "message": f"Expiry date {text!r} is not recognised."}],
            )
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _parse_valid_days(value: Any) -> tuple[list[int], str | None]:
    text = _cell_text(value).lower()
    if not text or "varies by card" in text:
        return [], "Valid weekdays need review."
    if "all days" in text or re.search(r"\bany\b", text):
        return _ALL_DAYS.copy(), None
    if "weekend" in text:
        return [5, 6], None
    range_match = re.search(
        r"\b(sun(?:day)?|mon(?:day)?|tue(?:s|sday)?|wed(?:nesday)?|"
        r"thu(?:r|rs|rsday|rsday)?|fri(?:day)?|sat(?:urday)?)\s+to\s+"
        r"(sun(?:day)?|mon(?:day)?|tue(?:s|sday)?|wed(?:nesday)?|"
        r"thu(?:r|rs|rsday|rsday)?|fri(?:day)?|sat(?:urday)?)\b",
        text,
    )
    if range_match:
        start = _DAY_NAMES[range_match.group(1)]
        end = _DAY_NAMES[range_match.group(2)]
        days = [start]
        while days[-1] != end:
            days.append((days[-1] + 1) % 7)
        return sorted(days), None
    days = {
        day
        for name, day in _DAY_NAMES.items()
        if re.search(rf"\b{re.escape(name)}\b", text)
    }
    if days:
        return sorted(days), None
    return [], "Valid weekdays need review."


def _card_types(card_type: Any, card_name: Any) -> tuple[str, ...]:
    raw_type = _cell_text(card_type).lower()
    card_description = _cell_text(card_name).lower()
    if "upi" in raw_type or "wallet" in raw_type or "platform" in raw_type or "pay later" in raw_type:
        return ()
    has_credit = "credit" in raw_type or "credit" in card_description
    has_debit = "debit" in raw_type or "debit" in card_description
    if has_credit and has_debit:
        return ("CREDIT", "DEBIT")
    if has_credit:
        return ("CREDIT",)
    if has_debit:
        return ("DEBIT",)
    return ()


def _category(value: Any, merchant: str) -> str:
    source = _cell_text(value).upper()
    combined = f"{source} {_cell_text(merchant).upper()}"
    if "MOVIE" in combined or "CINEMA" in combined:
        return "MOVIES"
    if "GROCERY" in combined or "INSTAMART" in combined or "ZEPTO" in combined:
        return "GROCERY"
    if "DINING" in combined or "RESTAURANT" in combined or "FOOD" in combined:
        return "DINING"
    if "SHOPPING" in combined or "ELECTRONICS" in combined:
        return "SHOPPING"
    if "HOTEL" in combined:
        return "HOTELS"
    if "TRAVEL" in combined or "FLIGHT" in combined:
        return "TRAVEL"
    return source or "OTHER"


def _normalised_identity(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def _network(bank_name: str, card_name: str) -> str:
    text = f"{bank_name} {card_name}".lower()
    if "mastercard" in text or "master card" in text:
        return "MASTERCARD"
    for network in ("visa", "rupay", "amex", "diners"):
        if network in text:
            return network.upper()
    return "OTHER"


def parse_offer_workbook(content: bytes) -> list[WorkbookOfferRow]:
    if not content:
        raise OfferWorkbookError("The selected workbook is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise OfferWorkbookError("The workbook exceeds the 10 MB upload limit.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_WORKBOOK_ENTRIES:
                raise OfferWorkbookError("The workbook contains too many internal files.")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise OfferWorkbookError("The workbook expands beyond the 50 MB safety limit.")
    except BadZipFile as exc:
        raise OfferWorkbookError("The selected file is not a valid .xlsx workbook.") from exc

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise OfferWorkbookError("The selected file is not a readable .xlsx workbook.") from exc

    imported_rows: list[WorkbookOfferRow] = []
    errors: list[dict[str, Any]] = []
    processed_rows = 0
    try:
        worksheets = workbook.worksheets
        if not worksheets:
            raise OfferWorkbookError("The workbook has no worksheets.")
        for sheet in worksheets:
            header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers: dict[int, str] = {}
            normalised_headers: set[str] = set()
            for column, value in enumerate(header_values):
                name = _cell_text(value)
                if not name:
                    continue
                normalised = _normalise_header(name)
                if normalised in normalised_headers:
                    errors.append({
                        "sheet": sheet.title,
                        "row": 1,
                        "message": f"Duplicate column header: {name}.",
                    })
                normalised_headers.add(normalised)
                headers[column] = name
            header_lookup = {
                _normalise_header(name): column for column, name in headers.items()
            }
            if "merchant" not in header_lookup or "offerheadline" not in header_lookup:
                errors.append({
                    "sheet": sheet.title,
                    "row": 1,
                    "message": "Each sheet must include Merchant and Offer Headline columns.",
                })
                continue

            for source_row, cells in enumerate(
                sheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                values = {
                    column: value
                    for column, value in enumerate(cells)
                    if value is not None and _cell_text(value)
                }
                if not values:
                    continue
                processed_rows += 1
                if processed_rows > MAX_OFFER_ROWS:
                    errors.append({
                        "sheet": sheet.title,
                        "row": source_row,
                        "message": f"The workbook exceeds the {MAX_OFFER_ROWS:,}-row limit.",
                    })
                    break
                unnamed = sorted(column + 1 for column in values if column not in headers)
                if unnamed:
                    errors.append({
                        "sheet": sheet.title,
                        "row": source_row,
                        "message": f"Data is present in column(s) without a header: {unnamed}.",
                    })
                    continue
                source_data = {
                    name: _source_value(values.get(column))
                    for column, name in headers.items()
                }
                row_lookup = {
                    _normalise_header(headers[column]): value
                    for column, value in values.items()
                }
                merchant = _cell_text(row_lookup.get("merchant"))
                headline = _cell_text(row_lookup.get("offerheadline"))
                if not merchant or not headline:
                    errors.append({
                        "sheet": sheet.title,
                        "row": source_row,
                        "message": "Merchant and Offer Headline must not be blank.",
                    })
                    continue

                try:
                    min_spend_value = next(
                        (
                            value
                            for key, value in row_lookup.items()
                            if key in {"minpurchase", "minimumspend", "minspend"}
                        ),
                        None,
                    )
                    max_discount_value = next(
                        (
                            value
                            for key, value in row_lookup.items()
                            if key in {"maxdiscountinr", "maximumdiscountinr", "maxdiscount"}
                        ),
                        None,
                    )
                    expires_value = next(
                        (value for key, value in row_lookup.items() if "expirydate" in key),
                        None,
                    )
                    valid_value = next(
                        (value for key, value in row_lookup.items() if key == "validdays"),
                        None,
                    )
                    bank_value = next(
                        (
                            value
                            for key, value in row_lookup.items()
                            if key in {"bankname", "bankprovider", "bank"}
                        ),
                        None,
                    )
                    card_value = next(
                        (value for key, value in row_lookup.items() if key == "cardname"),
                        None,
                    )
                    type_value = next(
                        (value for key, value in row_lookup.items() if key == "cardtype"),
                        None,
                    )
                    min_spend = _number(
                        min_spend_value,
                        field="Minimum purchase",
                        sheet=sheet.title,
                        row=source_row,
                    )
                    max_discount = _number(
                        max_discount_value,
                        field="Maximum discount",
                        sheet=sheet.title,
                        row=source_row,
                    )
                    expires_on = _parse_expiry(
                        expires_value,
                        sheet=sheet.title,
                        row=source_row,
                    )
                    valid_days, schedule_error = _parse_valid_days(valid_value)
                except OfferWorkbookError as exc:
                    errors.extend(exc.errors)
                    continue

                card_types = _card_types(type_value, card_value)
                reason = schedule_error
                if not card_types:
                    reason = "Payment type is not a supported bank card."
                imported_rows.append(WorkbookOfferRow(
                    source_sheet=sheet.title,
                    source_row=source_row,
                    source_data=source_data,
                    merchant_name=merchant,
                    category=_category(row_lookup.get("category"), merchant),
                    offer_headline=headline,
                    coupon_code=_cell_text(row_lookup.get("couponcode")) or None,
                    min_spend=min_spend or 0,
                    max_discount=max_discount,
                    valid_days=valid_days,
                    expires_on=expires_on,
                    bank_name=_cell_text(bank_value) or None,
                    card_name=_cell_text(card_value) or None,
                    card_types=card_types,
                    is_matchable=bool(card_types and valid_days),
                    unmatched_reason=reason,
                ))
    finally:
        workbook.close()

    if errors:
        raise OfferWorkbookError(
            f"Workbook validation failed for {len(errors)} row(s); nothing was imported.",
            errors,
        )
    if not imported_rows:
        raise OfferWorkbookError("The workbook contains no offer rows.")
    return imported_rows


def import_offer_rows(
    rows: list[WorkbookOfferRow],
    *,
    source_file: str,
    source_hash: str,
) -> dict[str, Any]:
    expected_keys = {(row.source_sheet, row.source_row) for row in rows}
    existing_rows = {
        (document["source_sheet"], document["source_row"])
        for document in merchant_offers.find(
            {"source_hash": source_hash},
            {"_id": 0, "source_sheet": 1, "source_row": 1},
        )
    }
    if existing_rows == expected_keys:
        raise OfferWorkbookError(
            "This exact workbook has already been uploaded.",
            [{"message": "No rows were imported because the workbook is already present."}],
            status_code=409,
        )
    if not existing_rows.issubset(expected_keys):
        raise OfferWorkbookError(
            "This workbook hash is already associated with a different row set.",
            [{"message": "No rows were imported because the source rows do not match."}],
            status_code=409,
        )

    cards_by_identity: dict[tuple[str, str, str], dict[str, Any]] = {}
    for card in master_cards.find({}, {"_id": 0}):
        identity = (
            _normalised_identity(card["bank_name"]),
            _normalised_identity(card["card_name"]),
            card["card_type"],
        )
        cards_by_identity[identity] = card

    cards_to_create: dict[tuple[str, str, str], dict[str, Any]] = {}
    identities_by_row: dict[tuple[str, int], list[tuple[str, str, str]]] = {}
    for row in rows:
        row_key = (row.source_sheet, row.source_row)
        identities: list[tuple[str, str, str]] = []
        if row.bank_name and row.card_name:
            for card_type in row.card_types:
                if not _cell_text(row.bank_name):
                    continue
                identity = (
                    _normalised_identity(row.bank_name),
                    _normalised_identity(row.card_name),
                    card_type,
                )
                identities.append(identity)
                if identity not in cards_by_identity and identity not in cards_to_create:
                    cards_to_create[identity] = {
                        "bank_name": row.bank_name,
                        "card_name": row.card_name,
                        "card_type": card_type,
                        "network": _network(row.bank_name, row.card_name),
                    }
        identities_by_row[row_key] = identities

    for identity, card in cards_to_create.items():
        card["id"] = new_catalogue_id("master_cards", master_cards)
        master_cards.insert_one(card)
        cards_by_identity[identity] = card

    write_rows = [row for row in rows if (row.source_sheet, row.source_row) not in existing_rows]
    offer_documents = []
    document_matchability = Counter()
    for row in write_rows:
        row_key = (row.source_sheet, row.source_row)
        card_ids = [
            cards_by_identity[identity]["id"]
            for identity in identities_by_row[row_key]
        ]
        is_matchable = row.is_matchable and bool(card_ids)
        document_matchability["matchable" if is_matchable else "unmatchable"] += 1
        offer_documents.append({
            "id": new_catalogue_id("merchant_offers", merchant_offers),
            "merchant_name": row.merchant_name,
            "category": row.category,
            "offer_headline": row.offer_headline,
            "coupon_code": row.coupon_code,
            "min_spend": row.min_spend,
            "max_discount": row.max_discount,
            "valid_days": row.valid_days,
            "applicable_master_card_ids": card_ids,
            "expires_on": row.expires_on,
            "is_matchable": is_matchable,
            "unmatched_reason": row.unmatched_reason if not is_matchable else None,
            "source_file": source_file,
            "source_hash": source_hash,
            "source_sheet": row.source_sheet,
            "source_row": row.source_row,
            "source_data": row.source_data,
            "imported_at": datetime.now(timezone.utc),
        })

    inserted_count = 0
    bulk_write_errors: list[dict[str, Any]] = []
    if offer_documents:
        try:
            result = merchant_offers.insert_many(offer_documents, ordered=True)
            inserted_count = len(result.inserted_ids)
        except BulkWriteError as exc:
            inserted_count = exc.details.get("nInserted", 0)
            bulk_write_errors = [
                {
                    "sheet": offer_documents[error["index"]]["source_sheet"],
                    "row": offer_documents[error["index"]]["source_row"],
                    "message": f"MongoDB write failed (code {error.get('code')}).",
                }
                for error in exc.details.get("writeErrors", [])
            ]
        except PyMongoError as exc:
            raise OfferWorkbookError(
                "MongoDB could not save the uploaded offers.",
                [{"message": f"Database write failed ({type(exc).__name__})."}],
                status_code=500,
            ) from exc

    stored_keys = {
        (document["source_sheet"], document["source_row"])
        for document in merchant_offers.find(
            {"source_hash": source_hash},
            {"_id": 0, "source_sheet": 1, "source_row": 1},
        )
    }
    missing_keys = sorted(expected_keys - stored_keys)
    persisted = [row for row in rows if (row.source_sheet, row.source_row) in stored_keys]
    return {
        "source_file": source_file,
        "rows_total": len(rows),
        "rows_imported": len(stored_keys & expected_keys),
        "rows_inserted_this_request": inserted_count,
        "cards_created": len(cards_to_create),
        "card_matchable_rows": sum(row.is_matchable for row in persisted),
        "unmatched_rows": sum(not row.is_matchable for row in persisted),
        "rows_without_expiry": sum(row.expires_on is None for row in persisted),
        "rows_without_valid_days": sum(not row.valid_days for row in persisted),
        "sheets": dict(Counter(row.source_sheet for row in rows)),
        "missing_rows": [
            {"sheet": sheet, "row": row} for sheet, row in missing_keys
        ],
        "errors": bulk_write_errors,
    }
